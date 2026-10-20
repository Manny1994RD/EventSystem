import datetime
import shutil
import tempfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from accounts.models import Role, User

from . import services
from .models import (
    ChangeRequest, Church, Event, EventReport, EventStatus, EventType, FlyerStatus, Notification, SiteSettings,
)

MEDIA = tempfile.mkdtemp()
PNG = b"\x89PNG\r\n\x1a\n" + b"0" * 100


@override_settings(MEDIA_ROOT=MEDIA, SECURE_PROXY_SSL_HEADER=None, SESSION_COOKIE_SECURE=False,
                   CSRF_COOKIE_SECURE=False, PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class Base(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    def setUp(self):
        self.santiago = Church.objects.create(name="Santiago")
        self.vega = Church.objects.create(name="La Vega")
        self.moca = Church.objects.create(name="Moca")
        self.type = EventType.objects.create(name_es="Seminario", name_en="Seminar")
        self.rep = self.make("rep", Role.CHURCH_REP, self.santiago)
        self.leader_vega = self.make("leader", Role.CHURCH_LEADER, self.vega)
        self.rep_moca = self.make("moca", Role.CHURCH_REP, self.moca)
        self.regional = self.make("regional", Role.REGIONAL_REP, admin=True)
        self.overseer = self.make("overseer", Role.OVERSEER)
        self.designer = self.make("designer", Role.DESIGNER)
        self.day = timezone.localdate() + datetime.timedelta(days=40)

    def make(self, username, role, church=None, admin=False):
        return User.objects.create_user(username=username, password="pass12345!", role=role, church=church,
                                        is_app_admin=admin, first_name=username)

    def event(self, church, user, date=None, submit=False):
        e = Event.objects.create(
            church=church, title=f"Event {church}", event_type=self.type, date=date or self.day,
            start_time=datetime.time(18), location="Here", description="x", contact_name="c",
            contact_phone="1", created_by=user,
        )
        if submit:
            services.submit_event(e, user)
        return e


class DailyLimitTests(Base):
    def test_third_event_on_same_day_is_blocked(self):
        self.event(self.santiago, self.rep, submit=True)
        self.event(self.vega, self.leader_vega, submit=True)
        third = self.event(self.moca, self.rep_moca)
        with self.assertRaises(services.DayFull):
            services.submit_event(third, self.rep_moca)
        third.refresh_from_db()
        self.assertEqual(third.status, EventStatus.DRAFT)

    def test_exception_request_then_reviewer_override(self):
        self.event(self.santiago, self.rep, submit=True)
        self.event(self.vega, self.leader_vega, submit=True)
        third = self.event(self.moca, self.rep_moca)
        services.submit_event(third, self.rep_moca, exception_reason="Visiting preacher")
        self.assertTrue(third.exception_requested)
        with self.assertRaises(services.DayFull):
            services.review_event(third, self.regional, "approve")
        services.review_event(third, self.regional, "approve", override=True)
        third.refresh_from_db()
        self.assertEqual(third.status, EventStatus.APPROVED)
        self.assertTrue(third.limit_override)
        self.assertEqual(third.override_by, self.regional)

    def test_drafts_and_declined_do_not_take_a_spot(self):
        self.event(self.santiago, self.rep)  # draft
        declined = self.event(self.vega, self.leader_vega, submit=True)
        services.review_event(declined, self.overseer, "decline", "no")
        self.event(self.moca, self.rep_moca, submit=True)
        self.assertFalse(services.day_is_full(self.day))

    def test_limit_is_a_setting(self):
        s = SiteSettings.load()
        s.max_events_per_day = 1
        s.save()
        self.event(self.santiago, self.rep, submit=True)
        with self.assertRaises(services.DayFull):
            services.submit_event(self.event(self.vega, self.leader_vega), self.leader_vega)


class WorkflowTests(Base):
    def test_church_leader_or_rep_can_submit_and_reviewers_are_notified(self):
        e = self.event(self.vega, self.leader_vega, submit=True)
        self.assertEqual(e.status, EventStatus.SUBMITTED)
        notified = set(Notification.objects.filter(event=e).values_list("user__username", flat=True))
        self.assertEqual(notified, {"regional", "overseer"})

    def test_either_reviewer_can_approve(self):
        a = self.event(self.santiago, self.rep, submit=True)
        b = self.event(self.vega, self.leader_vega, date=self.day + datetime.timedelta(days=1), submit=True)
        services.review_event(a, self.regional, "approve")
        services.review_event(b, self.overseer, "approve")
        self.assertEqual(Event.objects.filter(status=EventStatus.APPROVED).count(), 2)

    def test_church_cannot_approve(self):
        e = self.event(self.santiago, self.rep, submit=True)
        with self.assertRaises(services.NotAllowed):
            services.review_event(e, self.rep, "approve")

    def test_late_submission_is_flagged(self):
        e = self.event(self.santiago, self.rep)
        q = e.quarter
        q.submission_deadline = timezone.localdate() - datetime.timedelta(days=1)
        q.save()
        services.submit_event(e, self.rep)
        self.assertTrue(e.is_late)

    def test_reviewer_moving_date_needs_no_reapproval(self):
        e = self.event(self.santiago, self.rep, submit=True)
        services.review_event(e, self.regional, "approve")
        old = {"date": e.date, "start_time": e.start_time, "location": e.location}
        e.date = e.date + datetime.timedelta(days=3)
        services.reviewer_update(e, self.overseer, old)
        e.refresh_from_db()
        self.assertEqual(e.status, EventStatus.APPROVED)
        self.assertTrue(Notification.objects.filter(user=self.rep, event=e, text_en__contains="updated").exists())

    def test_church_change_request_flow(self):
        e = self.event(self.santiago, self.rep, submit=True)
        services.review_event(e, self.regional, "approve")
        new_date = self.day + datetime.timedelta(days=5)
        services.request_change(e, self.rep, ChangeRequest(new_date=new_date, reason="Rain"))
        e.refresh_from_db()
        self.assertEqual(e.date, self.day)  # unchanged until approved
        self.assertTrue(Notification.objects.filter(user=self.overseer, text_en__contains="asks to change").exists())
        services.decide_change(e.pending_change, self.overseer, approve=True)
        e.refresh_from_db()
        self.assertEqual(e.date, new_date)
        self.assertEqual(e.status, EventStatus.APPROVED)


class FlyerTests(Base):
    def upload(self, event, user):
        return services.upload_flyer(event, user, SimpleUploadedFile("flyer.png", PNG, content_type="image/png"))

    def test_overseer_sends_to_designer_then_gives_final_ok(self):
        e = self.event(self.santiago, self.rep, submit=True)
        self.upload(e, self.rep)
        self.assertEqual(e.flyer_status, FlyerStatus.PENDING)
        services.review_flyer(e, self.overseer, approve=False, notes="Fix the date")
        self.assertEqual(e.flyer_status, FlyerStatus.DESIGNER)
        self.assertTrue(Notification.objects.filter(user=self.designer, event=e).exists())
        self.upload(e, self.designer)
        self.assertEqual(e.flyer_status, FlyerStatus.PENDING)
        self.assertEqual(e.latest_flyer.version, 2)
        services.review_flyer(e, self.overseer, approve=True)
        self.assertEqual(e.flyer_status, FlyerStatus.APPROVED)

    def test_designer_fix_auto_approved_when_setting_on(self):
        s = SiteSettings.load()
        s.auto_approve_designer_flyers = True
        s.save()
        e = self.event(self.santiago, self.rep, submit=True)
        self.upload(e, self.rep)
        services.review_flyer(e, self.overseer, approve=False, notes="Fix")
        self.upload(e, self.designer)
        self.assertEqual(e.flyer_status, FlyerStatus.APPROVED)

    def test_regional_rep_or_overseer_review_flyers(self):
        e = self.event(self.santiago, self.rep, submit=True)
        self.upload(e, self.rep)
        self.assertTrue(Notification.objects.filter(user=self.regional, text_en__contains="New flyer").exists())
        for user in (self.rep, self.designer):
            with self.assertRaises(services.NotAllowed):
                services.review_flyer(e, user, approve=True)
        services.review_flyer(e, self.regional, approve=True)
        self.assertEqual(e.flyer_status, FlyerStatus.APPROVED)

    def test_flyer_alerts(self):
        e = self.event(self.santiago, self.rep, date=timezone.localdate() + datetime.timedelta(days=10),
                       submit=True)
        self.assertEqual(e.flyer_alert, "due")
        e.date = timezone.localdate() + datetime.timedelta(days=5)
        self.assertEqual(e.flyer_alert, "overdue")

    def test_private_flyer_is_not_public(self):
        e = self.event(self.santiago, self.rep, submit=True)
        f = self.upload(e, self.rep)
        self.assertEqual(self.client.get(reverse("events:flyer_file", args=[f.pk])).status_code, 404)
        self.client.force_login(self.rep_moca)
        self.assertEqual(self.client.get(reverse("events:flyer_file", args=[f.pk])).status_code, 404)
        self.client.force_login(self.rep)
        self.assertEqual(self.client.get(reverse("events:flyer_file", args=[f.pk])).status_code, 200)


class PageTests(Base):
    """Every page opens for every role, in both languages, and church users only see their church."""

    def test_pages_render_for_each_role(self):
        e = self.event(self.santiago, self.rep, submit=True)
        services.upload_flyer(e, self.rep, SimpleUploadedFile("f.png", PNG))
        pages = ["events:home", "events:list", "events:calendar", "events:notifications", "profile"]
        for user in (self.rep, self.leader_vega, self.regional, self.overseer, self.designer):
            self.client.force_login(user)
            for lang in ("es", "en"):
                user.language = lang
                user.save()
                for name in pages:
                    self.assertEqual(self.client.get(reverse(name)).status_code, 200, (user, name))
        self.client.force_login(self.regional)
        for name in ("events:review_queue", "events:regional", "events:reports", "events:settings", "users",
                     "user_new", "events:report_excel", "events:report_pdf"):
            self.assertEqual(self.client.get(reverse(name)).status_code, 200, name)
        self.assertEqual(self.client.get(reverse("events:detail", args=[e.pk])).status_code, 200)
        self.client.force_login(self.overseer)
        self.assertEqual(self.client.get(reverse("events:flyer_queue")).status_code, 200)
        self.client.force_login(self.designer)
        self.assertEqual(self.client.get(reverse("events:designer_queue")).status_code, 200)

    def test_church_user_cannot_open_other_church_event(self):
        e = self.event(self.santiago, self.rep, submit=True)
        self.client.force_login(self.leader_vega)
        self.assertEqual(self.client.get(reverse("events:detail", args=[e.pk])).status_code, 403)
        self.assertEqual(self.client.get(reverse("events:regional")).status_code, 403)
        self.assertEqual(self.client.get(reverse("users")).status_code, 403)

    def test_public_event_hides_contact_phone(self):
        e = self.event(self.santiago, self.rep, submit=True)
        e.contact_phone = "809-555-7777"
        e.save()
        services.review_event(e, self.regional, "approve")
        resp = self.client.get(reverse("events:public_event", args=[e.pk]))
        self.assertContains(resp, e.title)
        self.assertNotContains(resp, "809-555-7777")

    def test_public_calendar_shows_only_approved(self):
        a = self.event(self.santiago, self.rep, submit=True)
        services.review_event(a, self.regional, "approve")
        b = self.event(self.vega, self.leader_vega, submit=True)
        b.title = "Secret pending"
        b.save()
        resp = self.client.get(reverse("events:public_calendar"), {"y": self.day.year, "m": self.day.month})
        self.assertContains(resp, a.title)
        self.assertNotContains(resp, "Secret pending")

    def test_create_and_submit_through_the_form(self):
        self.client.force_login(self.rep)
        resp = self.client.post(reverse("events:create"), {
            "title": "Culto", "event_type": self.type.pk, "date": self.day.isoformat(), "start_time": "18:00",
            "location": "Templo", "description": "x", "contact_name": "Ana", "contact_phone": "809",
        })
        e = Event.objects.get(title="Culto")
        self.assertEqual(e.church, self.santiago)
        self.assertRedirects(resp, reverse("events:detail", args=[e.pk]))
        self.client.post(reverse("events:submit", args=[e.pk]))
        e.refresh_from_db()
        self.assertEqual(e.status, EventStatus.SUBMITTED)

    def test_admin_creates_user(self):
        self.client.force_login(self.regional)
        resp = self.client.post(reverse("user_new"), {
            "first_name": "Maria", "username": "maria", "role": Role.CHURCH_REP, "church": self.moca.pk,
            "language": "es", "is_active": "on", "password": "una-clave-segura",
        })
        self.assertRedirects(resp, reverse("users"))
        self.assertTrue(User.objects.get(username="maria").check_password("una-clave-segura"))

    def test_church_role_requires_church(self):
        self.client.force_login(self.regional)
        resp = self.client.post(reverse("user_new"), {
            "first_name": "X", "username": "x", "role": Role.CHURCH_REP, "language": "es",
            "is_active": "on", "password": "una-clave-segura",
        })
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(User.objects.filter(username="x").exists())


class EmailTests(Base):
    def turn_on(self, **extra):
        s = SiteSettings.load()
        s.email_enabled = True
        s.email_from = "regional@example.com"
        for k, v in extra.items():
            setattr(s, k, v)
        s.save()

    def test_email_is_off_by_default(self):
        from django.core import mail
        from django.core.management import call_command

        self.rep.email = "rep@example.com"
        self.rep.save()
        self.event(self.santiago, self.rep, submit=True)
        call_command("send_reminders", stdout=open("/dev/null", "w"))
        self.assertEqual(len(mail.outbox), 0)

    def test_notifications_are_emailed_in_each_persons_language(self):
        from django.core import mail

        self.turn_on(email_notifications=True)
        self.regional.email, self.regional.language = "r@example.com", "en"
        self.regional.save()
        self.overseer.email, self.overseer.language = "o@example.com", "es"
        self.overseer.save()
        with self.captureOnCommitCallbacks(execute=True):
            self.event(self.santiago, self.rep, submit=True)
        bodies = {m.to[0]: m.body for m in mail.outbox}
        self.assertIn("sent", bodies["r@example.com"])
        self.assertIn("envió", bodies["o@example.com"])

    def test_quarter_deadline_reminder_sent_once(self):
        from django.core import mail
        from django.core.management import call_command

        from .models import Quarter

        self.turn_on()
        self.rep.email = "rep@example.com"
        self.rep.save()
        q = Quarter.current().next()
        q.submission_deadline = timezone.localdate() + datetime.timedelta(days=7)
        q.save()
        call_command("send_reminders", stdout=open("/dev/null", "w"))
        call_command("send_reminders", stdout=open("/dev/null", "w"))
        self.assertEqual([m.to for m in mail.outbox], [["rep@example.com"]])
        self.assertIn("Santiago", mail.outbox[0].body)

    def test_flyer_due_reminder(self):
        from django.core import mail
        from django.core.management import call_command

        self.turn_on()
        self.rep.email = "rep@example.com"
        self.rep.save()
        self.event(self.santiago, self.rep, date=timezone.localdate() + datetime.timedelta(days=14), submit=True)
        call_command("send_reminders", stdout=open("/dev/null", "w"))
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("flyer", mail.outbox[0].body.lower())

    def test_email_settings_page_keeps_saved_password(self):
        self.turn_on(smtp_password="secret-app-pass")
        self.client.force_login(self.regional)
        resp = self.client.post(reverse("events:email_settings"), {
            "email_enabled": "on", "email_from": "regional@example.com", "smtp_host": "smtp.gmail.com",
            "smtp_port": 587, "email_reminders": "on",
        })
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(SiteSettings.load().smtp_password, "secret-app-pass")
        self.assertNotContains(self.client.get(reverse("events:email_settings")), "secret-app-pass")


class PostEventReportTests(Base):
    def past_event(self):
        e = self.event(self.santiago, self.rep, submit=True)
        services.review_event(e, self.regional, "approve")
        Event.objects.filter(pk=e.pk).update(date=timezone.localdate() - datetime.timedelta(days=1))
        e.refresh_from_db()
        return e

    def post(self, e, **extra):
        data = {"members_attended": 40, "supporters_attended": 12, "supporting_churches": [self.vega.pk],
                "members_supporting": 8, "fruits": 3, "bible_studies": 5, "invites": 25, "cost": "3500.00",
                "fragrance": "A family accepted a Bible study", **extra}
        return self.client.post(reverse("events:report_edit", args=[e.pk]), data)

    def test_church_submits_report_with_photos_and_reviewers_are_notified(self):
        self.type.full_report = True
        self.type.save()
        e = self.past_event()
        self.client.force_login(self.rep)
        photos = [SimpleUploadedFile(f"p{i}.jpg", PNG, content_type="image/jpeg") for i in range(2)]
        resp = self.post(e, submit="1", photos=photos)
        self.assertRedirects(resp, reverse("events:detail", args=[e.pk]))
        r = e.report
        self.assertEqual(r.status, "submitted")
        self.assertEqual(r.total_attendance, 52)
        self.assertEqual(r.photos.count(), 2)
        self.assertEqual(list(r.supporting_churches.all()), [self.vega])
        self.assertEqual((r.members_supporting, r.fruits, r.bible_studies), (8, 3, 5))
        self.assertEqual(r.fragrance, "A family accepted a Bible study")
        for user in (self.regional, self.overseer):
            self.assertTrue(Notification.objects.filter(user=user, text_en__contains="sent the report").exists())
        # Locked after sending.
        self.assertRedirects(self.client.get(reverse("events:report_edit", args=[e.pk])),
                             reverse("events:detail", args=[e.pk]))

    def test_save_draft_then_reviewer_returns_for_corrections(self):
        e = self.past_event()
        self.client.force_login(self.rep)
        self.post(e, save="1")
        self.assertEqual(e.report.status, "draft")
        self.post(e, submit="1")
        e.report.refresh_from_db()
        self.client.force_login(self.overseer)
        self.client.post(reverse("events:report_return", args=[e.pk]), {"comment": "Fix the count"})
        e.report.refresh_from_db()
        self.assertEqual(e.report.status, "draft")
        self.assertTrue(Notification.objects.filter(user=self.rep, text_en__contains="corrections").exists())

    def test_no_report_before_event_or_for_other_church(self):
        future = self.event(self.santiago, self.rep, submit=True)
        services.review_event(future, self.regional, "approve")
        self.client.force_login(self.rep)
        self.assertEqual(self.client.get(reverse("events:report_edit", args=[future.pk])).status_code, 403)
        past = self.past_event()
        self.client.force_login(self.leader_vega)
        self.assertEqual(self.client.get(reverse("events:report_edit", args=[past.pk])).status_code, 403)

    def test_pages_and_excel_include_reports(self):
        e = self.past_event()
        self.client.force_login(self.rep)
        self.post(e, submit="1", photos=[SimpleUploadedFile("p.jpg", PNG)])
        photo = e.report.photos.get()
        self.assertEqual(self.client.get(reverse("events:post_reports")).status_code, 200)
        self.assertEqual(self.client.get(reverse("events:report_photo", args=[photo.pk])).status_code, 200)
        self.client.force_login(self.rep_moca)
        self.assertEqual(self.client.get(reverse("events:report_photo", args=[photo.pk])).status_code, 404)
        self.client.force_login(self.regional)
        resp = self.client.get(reverse("events:post_reports"), {"quarter": e.quarter_id})
        self.assertContains(resp, e.title)
        self.assertEqual(self.client.get(reverse("events:detail", args=[e.pk])).status_code, 200)
        from io import BytesIO

        from openpyxl import load_workbook

        resp = self.client.get(reverse("events:report_excel"), {"quarter": e.quarter_id})
        wb = load_workbook(BytesIO(resp.content))
        sheet = wb.worksheets[1]
        self.assertEqual(sheet.cell(row=2, column=6).value, 40)

    def test_basic_report_for_non_bible_social_events(self):
        cleanup = EventType.objects.create(name_es="Jornada de limpieza", name_en="Cleanup")
        e = self.past_event()
        Event.objects.filter(pk=e.pk).update(event_type=cleanup)
        self.client.force_login(self.rep)
        resp = self.client.get(reverse("events:report_edit", args=[e.pk]))
        self.assertNotContains(resp, 'name="fruits"')
        self.assertNotContains(resp, 'name="members_supporting"')
        self.assertContains(resp, 'name="invites"')
        self.assertContains(resp, 'name="fragrance"')
        self.post(e, submit="1")
        r = EventReport.objects.get(event=e)
        self.assertEqual((r.members_attended, r.invites, r.fruits), (40, 25, 0))

    def test_report_reminder_email(self):
        from django.core import mail
        from django.core.management import call_command

        s = SiteSettings.load()
        s.email_enabled, s.email_from = True, "r@example.com"
        s.save()
        self.rep.email = "rep@example.com"
        self.rep.save()
        self.past_event()
        call_command("send_reminders", stdout=open("/dev/null", "w"))
        self.assertTrue(any("report" in m.subject.lower() or "informe" in m.subject.lower() for m in mail.outbox))


class CalendarSlideshowTests(Base):
    def approved(self, church, user, days):
        e = self.event(church, user, date=timezone.localdate() + datetime.timedelta(days=days), submit=True)
        services.review_event(e, self.regional, "approve")
        return e

    def test_public_calendar_shows_upcoming_slides_with_whatsapp_invite(self):
        soon = self.approved(self.santiago, self.rep, 1)
        soon.title = "Culto de jóvenes"
        soon.save()
        later = self.approved(self.vega, self.leader_vega, 30)
        later.title = "Far away event"
        later.save()
        pending = self.event(self.moca, self.rep_moca, date=timezone.localdate() + datetime.timedelta(days=2),
                             submit=True)
        pending.title = "Not approved yet"
        pending.save()
        resp = self.client.get(reverse("events:public_calendar"))
        slides = resp.context["slides"]
        self.assertEqual([s["event"] for s in slides], [soon])
        self.assertContains(resp, "https://wa.me/?text=")
        self.assertIn("Culto", slides[0]["whatsapp"])
        self.assertIn(f"/public/event/{soon.pk}/", slides[0]["link"])

    def test_approved_flyer_shows_on_slide(self):
        e = self.approved(self.santiago, self.rep, 1)
        f = services.upload_flyer(e, self.rep, SimpleUploadedFile("f.png", PNG))
        services.review_flyer(e, self.overseer, approve=True)
        resp = self.client.get(reverse("events:public_calendar"))
        self.assertContains(resp, reverse("events:flyer_file", args=[f.pk]))
        self.client.force_login(self.rep_moca)
        self.assertEqual(self.client.get(reverse("events:calendar")).status_code, 200)


class EventTypeToggleTests(Base):
    def test_admin_toggles_full_report(self):
        self.client.force_login(self.regional)
        self.client.post(reverse("events:type_toggle_report", args=[self.type.pk]))
        self.type.refresh_from_db()
        self.assertTrue(self.type.full_report)
        self.client.post(reverse("events:type_toggle_report", args=[self.type.pk]))
        self.type.refresh_from_db()
        self.assertFalse(self.type.full_report)

    def test_church_user_cannot_toggle(self):
        self.client.force_login(self.rep)
        resp = self.client.post(reverse("events:type_toggle_report", args=[self.type.pk]))
        self.assertEqual(resp.status_code, 403)

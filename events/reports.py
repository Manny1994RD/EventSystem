"""Quarter report downloads (Excel and PDF) for the Overseer and the Regional Events Rep."""

import io

from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.utils.translation import gettext as _

from .forms import ReportFilterForm
from .models import Event, EventReport, EventStatus, Quarter


def _check(user):
    if not (user.is_reviewer or user.can_manage):
        raise PermissionDenied


def _filtered(request):
    form = ReportFilterForm(request.GET or None)
    qs = Event.objects.exclude(status=EventStatus.DRAFT).select_related(
        "church", "event_type", "quarter", "reviewed_by")
    quarter = None
    if form.is_valid():
        quarter = form.cleaned_data.get("quarter")
        if form.cleaned_data.get("church"):
            qs = qs.filter(church=form.cleaned_data["church"])
        if form.cleaned_data.get("event_type"):
            qs = qs.filter(event_type=form.cleaned_data["event_type"])
        if form.cleaned_data.get("status"):
            qs = qs.filter(status=form.cleaned_data["status"])
    if quarter is None:
        quarter = Quarter.current().next()
        form = ReportFilterForm(initial={"quarter": quarter.pk, **request.GET.dict()})
    qs = qs.filter(quarter=quarter).order_by("date", "start_time")
    return form, quarter, list(qs)


def _columns():
    return [_("Date"), _("Time"), _("Church"), _("Event"), _("Type"), _("Place"), _("Contact"), _("Phone"),
            _("Status"), _("Flyer"), _("Late"), _("Exception"), _("Approved / reviewed by"), _("Help needed")]


def _row(e):
    time = e.start_time.strftime("%H:%M") + (f"–{e.end_time:%H:%M}" if e.end_time else "")
    return [
        e.date.strftime("%d/%m/%Y"), time, e.church.name, e.title, str(e.event_type), e.location,
        e.contact_name, e.contact_phone, e.get_status_display(), e.get_flyer_status_display(),
        _("Yes") if e.is_late else "", _("Yes") if e.limit_override else "",
        str(e.reviewed_by or ""), e.support_notes if e.needs_support else "",
    ]


@login_required
def report_page(request):
    _check(request.user)
    form, quarter, events = _filtered(request)
    return render(request, "events/reports.html", {
        "form": form, "quarter": quarter, "events": events, "query": request.GET.urlencode()})


@login_required
def report_excel(request):
    _check(request.user)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    _form, quarter, events = _filtered(request)
    wb = Workbook()
    ws = wb.active
    ws.title = str(quarter)
    ws.append([_("Events report") + f" — {quarter.label}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([_("Downloaded") + f": {timezone.localtime():%d/%m/%Y %H:%M}"])
    ws.append([])
    ws.append(_columns())
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
    for e in events:
        ws.append(_row(e))
    widths = [12, 12, 16, 32, 18, 30, 20, 15, 22, 22, 8, 10, 22, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    _post_event_sheet(wb, events, Font, PatternFill, get_column_letter)
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(),
                        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="eventos-{quarter.year}-Q{quarter.number}.xlsx"'
    return resp


@login_required
def report_pdf(request):
    _check(request.user)
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    _form, quarter, events = _filtered(request)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter), leftMargin=0.4 * inch, rightMargin=0.4 * inch,
                            topMargin=0.5 * inch, bottomMargin=0.5 * inch, title=f"{quarter}")
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=7.5, leading=9)
    head = small.clone("head", textColor=colors.white, fontName="Helvetica-Bold")
    cols = [_("Date"), _("Time"), _("Church"), _("Event"), _("Type"), _("Place"), _("Contact"), _("Status"),
            _("Flyer")]
    data = [[Paragraph(c, head) for c in cols]]
    for e in events:
        r = _row(e)
        data.append([Paragraph(str(v), small) for v in
                     (r[0], r[1], r[2], r[3], r[4], r[5], f"{r[6]}<br/>{r[7]}", r[8], r[9])])
    table = Table(data, repeatRows=1, colWidths=[w * inch for w in [0.75, 0.75, 1.0, 1.9, 1.0, 1.7, 1.2, 1.0, 1.0]])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story = [
        Paragraph(_("Events report") + f" — {quarter.label}", styles["Title"]),
        Paragraph(_("Downloaded") + f": {timezone.localtime():%d/%m/%Y %H:%M} · "
                  + _("%(n)s events") % {"n": len(events)}, styles["Normal"]),
        Spacer(1, 10),
        table if events else Paragraph(_("No events match these filters."), styles["Normal"]),
    ]
    doc.build(story)
    resp = HttpResponse(buf.getvalue(), content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="eventos-{quarter.year}-Q{quarter.number}.pdf"'
    return resp


def _post_event_sheet(wb, events, Font, PatternFill, get_column_letter):
    ws = wb.create_sheet(_("After the event"))
    cols = [_("Date"), _("Church"), _("Event"), _("Report"), _("Members"), _("Supporters"),
            _("Supporting churches"), _("Served"), _("Fruits"), _("Bible studies"), _("Invites"), _("Cost (RD$)"),
            _("Photos"), _("Notes")]
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
    reports = {r.event_id: r for r in EventReport.objects.filter(event__in=events).prefetch_related(
        "supporting_churches", "photos")}
    numbers = ["members_attended", "supporters_attended", "people_served", "fruits", "bible_studies", "invites", "cost"]
    totals = dict.fromkeys(numbers, 0)
    for e in events:
        if e.status != EventStatus.APPROVED:
            continue
        r = reports.get(e.pk)
        if r is None or r.status != EventReport.Status.SUBMITTED:
            ws.append([e.date.strftime("%d/%m/%Y"), e.church.name, e.title,
                       _("Not sent") if r is None or r.status != EventReport.Status.SUBMITTED else ""])
            continue
        for n in numbers:
            totals[n] += getattr(r, n)
        ws.append([e.date.strftime("%d/%m/%Y"), e.church.name, e.title, _("Sent"), r.members_attended,
                   r.supporters_attended, ", ".join(c.name for c in r.supporting_churches.all()), r.people_served,
                   r.fruits, r.bible_studies, r.invites, float(r.cost), r.photos.count(), r.notes])
    ws.append([])
    ws.append([_("Total"), "", "", "", totals["members_attended"], totals["supporters_attended"], "",
               totals["people_served"], totals["fruits"], totals["bible_studies"], totals["invites"],
               float(totals["cost"])])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for i, w in enumerate([12, 16, 32, 12, 10, 11, 26, 10, 9, 13, 9, 12, 8, 50], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
